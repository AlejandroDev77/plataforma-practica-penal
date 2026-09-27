from __future__ import annotations

import io

import pymupdf
import pytest
from docx import Document
from httpx import ASGITransport, AsyncClient
from PIL import Image
from pydantic import SecretStr

from app.api.v1.routes import documents as documents_route
from app.core.config import Settings
from app.main import app
from app.modules.ingestion.service import DocumentExtractionError, DocumentExtractor


def make_pdf(*page_texts: str | None) -> bytes:
    document = pymupdf.open()
    for text in page_texts:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def make_docx() -> bytes:
    document = Document()
    document.add_paragraph("ACTA DE AUDIENCIA · 25 DE SEPTIEMBRE")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Defensa"
    table.cell(0, 1).text = "Solicitud de medidas cautelares"
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def make_image(image_format: str = "PNG", pages: int = 1) -> bytes:
    frames = [Image.new("RGB", (80, 40), "white") for _ in range(pages)]
    output = io.BytesIO()
    frames[0].save(output, format=image_format, save_all=True, append_images=frames[1:])
    for frame in frames:
        frame.close()
    return output.getvalue()


def test_pdf_text_is_extracted_with_real_page_locators() -> None:
    result = DocumentExtractor(Settings()).extract(
        "actuaciones.pdf",
        make_pdf("Auto interlocutorio sobre medidas cautelares"),
    )

    assert result.document_type == "pdf"
    assert result.page_count == 1
    assert result.pages[0].locator == "Página 1"
    assert "Auto interlocutorio" in result.pages[0].text
    assert result.pages[0].is_readable is True
    assert result.pages[0].used_ocr is False
    assert result.warnings == []


def test_docx_text_is_extracted_without_fabricating_physical_pagination() -> None:
    result = DocumentExtractor(Settings()).extract("actuaciones.docx", make_docx())

    assert result.document_type == "docx"
    assert result.page_count is None
    assert result.pages[0].locator == "Documento completo (paginación física no disponible)"
    assert "ACTA DE AUDIENCIA" in result.pages[0].text
    assert "Defensa | Solicitud de medidas cautelares" in result.pages[0].text
    assert result.warnings[0].code == "pagination_unavailable"


def test_scan_uses_ocr_only_when_native_pdf_text_is_below_threshold(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        DocumentExtractor,
        "_ocr_page",
        lambda self, page: ("Texto reconocido por OCR", None),
    )

    result = DocumentExtractor(Settings()).extract("escaneo.pdf", make_pdf(None))

    assert result.pages[0].text == "Texto reconocido por OCR"
    assert result.pages[0].used_ocr is True
    assert result.warnings == []


def test_scan_without_tesseract_is_reported_as_unreadable_not_as_success(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr("app.modules.ingestion.service.shutil.which", lambda _: None)

    result = DocumentExtractor(Settings()).extract("escaneo.pdf", make_pdf(None))

    assert result.pages[0].text == ""
    assert result.pages[0].is_readable is False
    assert result.warnings[0].code == "ocr_unavailable"


def test_multipage_tiff_keeps_page_boundaries(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        DocumentExtractor,
        "_ocr_page",
        lambda self, page: ("Contenido de la imagen", None),
    )

    result = DocumentExtractor(Settings()).extract("anexo.tiff", make_image("TIFF", pages=2))

    assert result.document_type == "image"
    assert result.page_count == 2
    assert [page.locator for page in result.pages] == ["Imagen 1", "Imagen 2"]
    assert all(page.used_ocr for page in result.pages)


def test_format_mismatch_is_rejected() -> None:
    extractor = DocumentExtractor(Settings())

    with pytest.raises(DocumentExtractionError) as error:
        extractor.extract("archivo.pdf", make_image())

    assert error.value.code == "invalid_pdf"
    assert error.value.status_code == 422


def test_empty_unsupported_and_oversized_inputs_are_rejected() -> None:
    extractor = DocumentExtractor(Settings(document_max_bytes=4))

    with pytest.raises(DocumentExtractionError) as empty:
        extractor.extract("vacío.pdf", b"")
    with pytest.raises(DocumentExtractionError) as unsupported:
        extractor.extract("presentación.pptx", b"dato")
    with pytest.raises(DocumentExtractionError) as oversized:
        extractor.extract("grande.pdf", b"%PDF-" + (b"0" * 10))

    assert empty.value.code == "empty_file"
    assert unsupported.value.code == "unsupported_format"
    assert oversized.value.code == "file_too_large"
    assert oversized.value.status_code == 413


async def test_internal_document_endpoint_requires_service_token(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        documents_route,
        "get_settings",
        lambda: Settings(service_token=SecretStr("local-internal-secret")),
    )
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post(
            "/api/v1/documents/extract",
            files={
                "file": ("actuaciones.pdf", make_pdf("Texto del expediente"), "application/pdf")
            },
        )

    assert response.status_code == 401
    assert response.json()["detail"]["code"] == "unauthorized"


async def test_internal_document_endpoint_returns_page_scoped_result(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        documents_route,
        "get_settings",
        lambda: Settings(service_token=SecretStr("local-internal-secret")),
    )
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        response = await client.post(
            "/api/v1/documents/extract",
            headers={"Authorization": "Bearer local-internal-secret"},
            files={
                "file": (
                    "actuaciones.pdf",
                    make_pdf("Texto completo de la actuación penal"),
                    "application/pdf",
                )
            },
        )

    assert response.status_code == 200
    payload = response.json()
    assert payload["document_type"] == "pdf"
    assert payload["pages"][0]["locator"] == "Página 1"
    assert "Texto completo de la actuación penal" in payload["pages"][0]["text"]
